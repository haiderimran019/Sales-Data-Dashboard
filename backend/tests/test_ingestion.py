from io import BytesIO
from pathlib import Path
from uuid import uuid4

import fitz
import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.services.ingestion.csv import CsvExtractor
from app.services.ingestion.detector import detect_file
from app.services.ingestion.docx import DocxExtractor
from app.services.ingestion.excel import ExcelExtractor
from app.services.ingestion.image import ImageExtractor
from app.services.ingestion.pdf import PdfExtractor
from app.services.ingestion.pptx import PptxExtractor
from app.services.ingestion.storage import LocalStorage


def write_csv(path: Path) -> None:
    path.write_text("name,value\nAlpha,1\nBeta,2\n", encoding="utf-8")


def write_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["name", "value"])
    sheet.append(["Alpha", 1])
    second = workbook.create_sheet("Notes")
    second.append(["note"])
    second.append(["hello"])
    workbook.save(path)


def write_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "A small report")
    document.save(path)
    document.close()


def write_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Report", level=1)
    document.add_paragraph("A paragraph")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "value"
    table.cell(1, 0).text = "one"
    document.save(path)


def write_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Report"
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "A slide"
    presentation.save(path)


def write_png(path: Path) -> None:
    Image.new("RGB", (8, 8), "white").save(path, format="PNG")


@pytest.mark.parametrize(
    ("extension", "writer", "expected"),
    [(".csv", write_csv, "csv"), (".xlsx", write_xlsx, "xlsx"), (".pdf", write_pdf, "pdf"), (".docx", write_docx, "docx"), (".pptx", write_pptx, "pptx"), (".png", write_png, "image")],
)
def test_supported_file_detection(tmp_path: Path, extension, writer, expected) -> None:
    path = tmp_path / f"sample{extension}"
    writer(path)

    detected = detect_file(path, path.name, None)

    assert detected.kind == expected


def test_unsupported_and_malformed_files_are_rejected(tmp_path: Path) -> None:
    unsupported = tmp_path / "script.exe"
    unsupported.write_bytes(b"MZ")
    malformed = tmp_path / "bad.pdf"
    malformed.write_bytes(b"not a pdf")

    with pytest.raises(ValueError, match="Unsupported"):
        detect_file(unsupported, unsupported.name)
    with pytest.raises(ValueError, match="signature"):
        detect_file(malformed, malformed.name)


def test_csv_extraction_is_bounded_and_profiles_values(tmp_path: Path) -> None:
    path = tmp_path / "sample.csv"
    path.write_text("name,value\nAlpha,1\nBeta,2\nGamma,3\n", encoding="utf-8")

    result = CsvExtractor().extract(path, detect_file(path, path.name), row_limit=2, text_limit=100)

    assert result.tables[0].row_count == 3
    assert len(result.tables[0].rows) == 2
    assert result.tables[0].data_types["value"] == "number"


def test_excel_preserves_multiple_sheets(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    write_xlsx(path)

    result = ExcelExtractor().extract(path, detect_file(path, path.name), row_limit=10, text_limit=100)

    assert [table.name for table in result.tables] == ["Sales", "Notes"]


def test_document_and_image_extractors_return_structured_previews(tmp_path: Path) -> None:
    pdf = tmp_path / "sample.pdf"
    docx = tmp_path / "sample.docx"
    pptx = tmp_path / "sample.pptx"
    image = tmp_path / "sample.png"
    write_pdf(pdf)
    write_docx(docx)
    write_pptx(pptx)
    write_png(image)

    assert PdfExtractor().extract(pdf, detect_file(pdf, pdf.name), row_limit=10, text_limit=100).metadata["page_count"] == 1
    assert DocxExtractor().extract(docx, detect_file(docx, docx.name), row_limit=10, text_limit=100).tables
    assert PptxExtractor().extract(pptx, detect_file(pptx, pptx.name), row_limit=10, text_limit=100).metadata["slide_count"] == 1
    assert ImageExtractor().extract(image, detect_file(image, image.name), row_limit=10, text_limit=100).metadata["ocr_status"] == "not_configured"


def test_storage_checksum_size_limit_and_path_traversal(tmp_path: Path) -> None:
    storage = LocalStorage(tmp_path / "storage")
    key = storage.key_for(uuid4(), uuid4(), ".csv")
    size, checksum = storage.save(BytesIO(b"name\nvalue\n"), key, max_bytes=100)

    assert size == 11
    assert len(checksum) == 64
    with pytest.raises(ValueError):
        storage.path_for("../outside")
    with pytest.raises(ValueError, match="size limit"):
        storage.save(BytesIO(b"0123456789"), key, max_bytes=4)
