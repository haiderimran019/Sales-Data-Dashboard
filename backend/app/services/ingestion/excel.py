from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.services.ingestion.base import Extractor
from app.services.ingestion.types import DetectedFile, NormalizedExtraction, TablePreview
from app.services.ingestion.csv import _unique_headers, _value_type


class ExcelExtractor(Extractor):
    kind = "xlsx"

    def extract(self, path: Path, detected: DetectedFile, *, row_limit: int, text_limit: int) -> NormalizedExtraction:
        tables: list[TablePreview] = []
        warnings: list[str] = []
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                first = next(rows, None)
                if not first or not any(value is not None and str(value).strip() for value in first):
                    warnings.append(f"Sheet '{sheet.title}' is blank and was skipped.")
                    continue
                headers = _unique_headers([str(value) if value is not None else "" for value in first])
                preview: list[dict[str, Any]] = []
                missing: dict[str, int] = {}
                types: dict[str, set[str]] = {}
                row_count = 0
                for raw_row in rows:
                    if not any(value is not None and str(value).strip() for value in raw_row):
                        continue
                    row_count += 1
                    values = list(raw_row) + [None] * max(0, len(headers) - len(raw_row))
                    record = {header: values[index] for index, header in enumerate(headers)}
                    for header, value in record.items():
                        if value is None or value == "":
                            missing[header] = missing.get(header, 0) + 1
                        types.setdefault(header, set()).add(_value_type(str(value) if value is not None else ""))
                    if len(preview) < row_limit:
                        preview.append(record)
                tables.append(TablePreview(name=sheet.title, columns=headers, rows=preview, row_count=row_count, missing_values=missing, data_types={key: (next(iter(value)) if len(value) == 1 else "mixed") for key, value in types.items()}))
        finally:
            workbook.close()
        if not tables:
            warnings.append("The workbook contains no non-empty sheets.")
        return NormalizedExtraction(detected_type="xlsx", tables=tables, metadata={"sheet_count": len(workbook.sheetnames)}, warnings=warnings, confidence=0.95)
