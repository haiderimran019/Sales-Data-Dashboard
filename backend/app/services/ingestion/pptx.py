from pathlib import Path

from pptx import Presentation

from app.services.ingestion.base import Extractor
from app.services.ingestion.types import DetectedFile, NormalizedExtraction, TablePreview, TextBlock


class PptxExtractor(Extractor):
    kind = "pptx"

    def extract(self, path: Path, detected: DetectedFile, *, row_limit: int, text_limit: int) -> NormalizedExtraction:
        presentation = Presentation(path)
        blocks: list[TextBlock] = []
        tables: list[TablePreview] = []
        warnings: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        blocks.append(TextBlock(location=f"slide {slide_number}", text=text[:text_limit]))
                if getattr(shape, "has_table", False):
                    table = shape.table
                    raw_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                    if raw_rows:
                        headers = [value or f"column_{i + 1}" for i, value in enumerate(raw_rows[0])]
                        tables.append(TablePreview(name=f"slide_{slide_number}_table", columns=headers, rows=[dict(zip(headers, row)) for row in raw_rows[1:row_limit + 1]], row_count=max(0, len(raw_rows) - 1)))
                if shape.shape_type == 13:
                    warnings.append(f"Slide {slide_number} contains an image; visual/OCR analysis is pending.")
        return NormalizedExtraction(detected_type="pptx", text_blocks=blocks, tables=tables, metadata={"slide_count": len(presentation.slides), "table_count": len(tables)}, warnings=warnings, confidence=0.85)
