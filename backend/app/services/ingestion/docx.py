from pathlib import Path

from docx import Document

from app.services.ingestion.base import Extractor
from app.services.ingestion.types import DetectedFile, NormalizedExtraction, TablePreview, TextBlock


class DocxExtractor(Extractor):
    kind = "docx"

    def extract(self, path: Path, detected: DetectedFile, *, row_limit: int, text_limit: int) -> NormalizedExtraction:
        document = Document(path)
        blocks: list[TextBlock] = []
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if text:
                blocks.append(TextBlock(location=f"paragraph {index + 1} ({paragraph.style.name})", text=text[:text_limit]))
        tables: list[TablePreview] = []
        for table_index, table in enumerate(document.tables, start=1):
            raw_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not raw_rows:
                continue
            headers = raw_rows[0] or [f"column_{i + 1}" for i in range(len(raw_rows[0]))]
            headers = [header or f"column_{i + 1}" for i, header in enumerate(headers)]
            rows = [dict(zip(headers, row)) for row in raw_rows[1:row_limit + 1]]
            tables.append(TablePreview(name=f"table_{table_index}", columns=headers, rows=rows, row_count=max(0, len(raw_rows) - 1)))
        return NormalizedExtraction(detected_type="docx", text_blocks=blocks, tables=tables, metadata={"paragraph_count": len(document.paragraphs), "table_count": len(tables)}, confidence=0.9)
